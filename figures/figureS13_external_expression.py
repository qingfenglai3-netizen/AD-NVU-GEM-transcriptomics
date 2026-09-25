import sys
from pathlib import Path
import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.gridspec import GridSpecFromSubplotSpec
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

base = r'<ANALYSIS_ROOT>/fuxian_test'
sea_dir = os.path.join(base, 'SEAAD_ADNC_vascular_GEM_validation')
bulk_dir = os.path.join(base, 'AD_bulk_GEO_GEM_external_screen')
outdir = r'<WORKDIR>/figures/S13_source'
os.makedirs(outdir, exist_ok=True)

sea_stats = pd.read_csv(os.path.join(sea_dir, 'SEAAD_ADNC_vascular_GEM_module_stats.csv'))
sea_pb = pd.read_csv(os.path.join(sea_dir, 'SEAAD_ADNC_vascular_sample_expression.csv'))
reg = pd.read_csv(os.path.join(bulk_dir, 'GSE5281_region_stratified_GEM_module_screen_v2.csv'))
gem_bulk = pd.read_csv(os.path.join(bulk_dir, 'AD_bulk_GEO_GEM_only_screen_results_v2_logscale_checked.csv'))

plt.rcParams.update({
    'font.family': 'Arial',
    'font.size': 6,
    'axes.titlesize': 6.5,
    'axes.labelsize': 6,
    'xtick.labelsize': 5.8,
    'ytick.labelsize': 5.8,
    'pdf.fonttype': 42,
    'ps.fonttype': 42,
})

fig = plt.figure(figsize=(6.69, 8.55), dpi=300)
gs = fig.add_gridspec(3, 1, height_ratios=[1.65, 1.25, 0.86], hspace=0.64)

sgA = GridSpecFromSubplotSpec(2, 2, subplot_spec=gs[0, 0], hspace=0.76, wspace=0.38)
plot_specs = [
    ('DLPFC_Endothelial', 'DLPFC Endothelial'),
    ('MTG_Endothelial', 'MTG Endothelial'),
    ('DLPFC_VLMC', 'DLPFC VLMC'),
    ('MTG_VLMC', 'MTG VLMC'),
]
colors = {'normal': '#4C78A8', 'dementia': '#D95F59'}
for k, (ds, title) in enumerate(plot_specs):
    ax = fig.add_subplot(sgA[k // 2, k % 2])
    d = sea_pb[(sea_pb.dataset == ds) & (sea_pb.gene == 'GEM') & (sea_pb.disease.isin(['normal', 'dementia']))].copy()
    groups = ['normal', 'dementia']
    data = [d.loc[d.disease == g, 'expr'].dropna().values for g in groups]
    bp = ax.boxplot(data, positions=[0, 1], widths=0.50, patch_artist=True, showfliers=False,
                    medianprops=dict(color='black', linewidth=0.8), whiskerprops=dict(linewidth=.65), capprops=dict(linewidth=.65))
    for patch, g in zip(bp['boxes'], groups):
        patch.set(facecolor=colors[g], alpha=.55, edgecolor='black', linewidth=.65)
    rng = np.random.default_rng(51 + k)
    for i, g in enumerate(groups):
        y = data[i]
        x = np.full(len(y), i) + rng.normal(0, 0.042, len(y))
        ax.scatter(x, y, s=5.5, color=colors[g], edgecolor='white', linewidth=.13, alpha=.9, zorder=3)
    st = sea_stats[(sea_stats.dataset == ds) & (sea_stats.gene == 'GEM') & (sea_stats.comparison == 'dementia_vs_normal')].iloc[0]
    ax.set_title(title, fontweight='bold', pad=2)
    ax.set_xticks([0, 1])
    ax.set_xticklabels(['Normal', 'Dementia'], rotation=18, ha='right', fontweight='bold')
    ax.set_ylabel('GEM expression' if k in [0, 2] else '', fontweight='bold')
    ax.text(0.04, .96, f"P={st.p_wilcox:.3g}; q={st.q_by_gene_comparison:.3g}", transform=ax.transAxes,
            ha='left', va='top', fontsize=5.8, fontweight='bold')
    ax.spines[['top', 'right']].set_visible(False)
fig.text(0.025, 0.958, 'A', fontsize=12, fontweight='bold')

axh = fig.add_subplot(gs[1, 0])
sel = ['GEM', 'NOTCH3', 'PDGFRB', 'PTH1R', 'PDE7B', 'SLC6A12', 'CACNA1C', 'RHBDF2']
heat = reg[reg.gene.isin(sel)].sort_values('P.Value').drop_duplicates(['subset', 'gene'])
mat = heat.pivot(index='gene', columns='subset', values='logFC').reindex(sel)
pv = heat.pivot(index='gene', columns='subset', values='P.Value').reindex(sel)
cols = [c for c in ['entorhinal cortex', 'hippocampus', 'medial temporal gyrus', 'posterior cingulate cortex', 'superior frontal gyrus', 'primary visual cortex'] if c in mat.columns]
mat = mat[cols]
pv = pv[cols]
vmax = max(np.nanmax(np.abs(mat.values)), 0.5)
im = axh.imshow(mat.values, cmap='RdBu_r', vmin=-vmax, vmax=vmax, aspect='auto')
axh.set_yticks(range(len(sel)))
axh.set_yticklabels(sel, fontweight='bold')
axh.set_xticks(range(len(cols)))
axh.set_xticklabels([c.replace(' cortex', '').replace(' gyrus', '') for c in cols], rotation=28, ha='right', fontweight='bold')
for i, g in enumerate(sel):
    for j, c in enumerate(cols):
        val = mat.loc[g, c]
        if pd.notna(val):
            star = '*' if pv.loc[g, c] < 0.05 else ''
            axh.text(j, i, f'{val:.2f}{star}', ha='center', va='center', fontsize=5.55, fontweight='bold')
axh.set_title('GSE5281 laser-captured neurons: regional AD-control logFC', fontweight='bold', pad=4)
cbar = fig.colorbar(im, ax=axh, fraction=.02, pad=.018)
cbar.ax.tick_params(labelsize=5.8)
cbar.set_label('logFC', fontsize=6, fontweight='bold')
fig.text(0.025, 0.465, 'B', fontsize=12, fontweight='bold')

axb = fig.add_subplot(gs[2, 0])
gem_region = heat[heat.gene == 'GEM'].set_index('subset').reindex(cols).reset_index()
overall = gem_bulk[(gem_bulk.GSE == 'GSE5281') & (gem_bulk.gene == 'GEM')].iloc[0]
x = np.arange(len(cols))
y = gem_region['logFC'].astype(float).values
p = gem_region['P.Value'].astype(float).values
bar_colors = ['#D95F59' if (yy > 0 and pp < 0.05) else '#F1A340' if yy > 0 else '#4C78A8' for yy, pp in zip(y, p)]
axb.bar(x, y, color=bar_colors, edgecolor='black', linewidth=.65, width=.60)
ymax = max(1.95, np.nanmax(y) + 0.42)
axb.set_ylim(0, ymax)
for i, (yy, pp) in enumerate(zip(y, p)):
    star = '*' if pp < 0.05 else ''
    axb.text(i, yy + 0.055, f'{yy:.2f}{star}', ha='center', va='bottom', fontsize=5.7, fontweight='bold')
axb.axhline(0, color='black', linewidth=.65)
axb.set_xticks(x)
axb.set_xticklabels([c.replace(' cortex', '').replace(' gyrus', '') for c in cols], rotation=25, ha='right', fontweight='bold')
axb.set_ylabel('GEM logFC', fontweight='bold')
axb.set_title('GSE5281 neuronal GEM by brain region', fontweight='bold', pad=7)
axb.spines[['top','right']].set_visible(False)
fig.text(0.025, 0.19, 'C', fontsize=12, fontweight='bold')

fig.subplots_adjust(top=.965, bottom=.105, left=.125, right=.945)
for ext, dpi in [('png', 600), ('tiff', 600), ('pdf', 600), ('svg', 600)]:
    fig.savefig(os.path.join(outdir, f'Supplementary_Figure_S13_external_validation_SEAAD_GSE5281_FINAL.{ext}'), dpi=dpi, bbox_inches='tight')
plt.close(fig)

print(outdir)
